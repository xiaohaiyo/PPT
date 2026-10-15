from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from .models import PresentationSpec, SlideSpec

W, H = Inches(13.333), Inches(7.5)
THEMES = {
    "corporate-blue": {"bg": "F5F8FC", "dark": "0B1F3A", "primary": "2457D6", "accent": "27B3A2", "text": "17243A"},
    "minimal-light": {"bg": "FAFAF8", "dark": "202020", "primary": "3A3A3A", "accent": "C9794A", "text": "252525"},
}


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _textbox(slide, x, y, w, h, text, size=24, color="17243A", bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True; frame.margin_left = frame.margin_right = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = text; paragraph.alignment = align
    paragraph.font.name = "Microsoft YaHei"; paragraph.font.size = Pt(size); paragraph.font.bold = bold; paragraph.font.color.rgb = _rgb(color)
    return shape


def _background(slide, color: str):
    fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = _rgb(color)


class PptxRenderer:
    def render(self, spec: PresentationSpec, output: Path) -> list[str]:
        prs = Presentation(); prs.slide_width = W; prs.slide_height = H
        theme = THEMES.get(spec.requirement.template_id, THEMES["corporate-blue"])
        warnings: list[str] = []
        for item in spec.slides:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            self._render_slide(slide, item, theme)
            notes = item.speaker_notes
            if item.citations:
                notes += "\n资料来源：" + "；".join(f"{x.document_name}" + (f" 第{x.page}页" if x.page else "") for x in item.citations)
            if notes.strip():
                slide.notes_slide.notes_text_frame.text = notes.strip()
            if len(item.title) > 30:
                warnings.append(f"第 {item.page} 页标题较长")
            if sum(len(x) for x in item.bullets) > 300:
                warnings.append(f"第 {item.page} 页正文可能过多")
        prs.save(output)
        return warnings

    def _render_slide(self, slide, item: SlideSpec, theme: dict[str, str]):
        if item.type == "cover":
            _background(slide, theme["dark"])
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.7), Inches(.85), Inches(.12), Inches(4.9))
            bar.fill.solid(); bar.fill.fore_color.rgb = _rgb(theme["accent"]); bar.line.fill.background()
            _textbox(slide, 1.2, 1.15, 10.8, 2.4, item.title, 38 if len(item.title) > 18 else 48, "FFFFFF", True)
            _textbox(slide, 1.25, 4.1, 9, .7, item.subtitle, 20, "C8D8EA")
            _textbox(slide, 1.25, 6.45, 5, .35, "AI WRITING ASSISTANT", 10, theme["accent"], True)
            return
        _background(slide, theme["bg"])
        _textbox(slide, .75, .42, 11.8, .72, item.title, 26 if len(item.title) > 22 else 30, theme["text"], True)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.75), Inches(1.28), Inches(1.05), Inches(.07))
        line.fill.solid(); line.fill.fore_color.rgb = _rgb(theme["accent"]); line.line.fill.background()
        if item.layout == "image-right-text-left" and item.image_path and Path(item.image_path).is_file():
            self._image(slide, item, theme)
        elif item.layout == "agenda-grid":
            self._agenda(slide, item, theme)
        elif item.layout in {"chart-left-text-right", "chart-right-text-left"} and item.chart and item.chart.categories and item.chart.values:
            self._chart(slide, item, theme)
        elif item.layout == "two-column":
            self._two_column(slide, item, theme)
        elif item.layout == "summary-cards":
            self._summary(slide, item, theme)
        elif item.layout == "section-divider":
            self._section(slide, item, theme)
        elif item.layout == "cards-grid":
            self._cards_grid(slide, item, theme)
        else:
            self._bullets(slide, item, theme)
        _textbox(slide, 11.8, 7.0, .75, .25, str(item.page).zfill(2), 9, "77869A", True, PP_ALIGN.RIGHT)
        if item.citations:
            source = "来源：" + "；".join(dict.fromkeys(x.document_name for x in item.citations))
            _textbox(slide, .78, 6.93, 9.8, .27, source, 8, "77869A")

    def _bullets(self, slide, item, theme):
        y = 1.72
        for index, bullet in enumerate(item.bullets[:5]):
            circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(.85), Inches(y + .08), Inches(.38), Inches(.38))
            circle.fill.solid(); circle.fill.fore_color.rgb = _rgb(theme["primary"]); circle.line.fill.background()
            _textbox(slide, .85, y + .08, .38, .38, str(index + 1), 10, "FFFFFF", True, PP_ALIGN.CENTER)
            _textbox(slide, 1.48, y, 10.7, .72, bullet, 19 if len(bullet) > 55 else 22, theme["text"])
            y += .94

    def _agenda(self, slide, item, theme):
        for i, value in enumerate(item.bullets[:8]):
            x = .85 + (i % 2) * 6.0; y = 1.65 + (i // 2) * 1.25
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.45), Inches(.92))
            card.fill.solid(); card.fill.fore_color.rgb = _rgb("FFFFFF"); card.line.color.rgb = _rgb("DDE5EF")
            _textbox(slide, x + .25, y + .12, .55, .65, str(i + 1).zfill(2), 14, theme["accent"], True)
            _textbox(slide, x + 1.0, y + .1, 4.1, .7, value, 16 if len(value) > 28 else 18, theme["text"], True)

    def _two_column(self, slide, item, theme):
        midpoint = max(1, (len(item.bullets) + 1) // 2)
        for col, values in enumerate((item.bullets[:midpoint], item.bullets[midpoint:])):
            x = .85 + col * 6.05
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.7), Inches(5.55), Inches(4.65))
            card.fill.solid(); card.fill.fore_color.rgb = _rgb("FFFFFF"); card.line.color.rgb = _rgb("DDE5EF")
            for i, value in enumerate(values[:4]): _textbox(slide, x + .4, 2.05 + i * .95, 4.75, .72, "•  " + value, 18, theme["text"])

    def _cards_grid(self, slide, item, theme):
        for index, value in enumerate(item.bullets[:4]):
            x = .85 + (index % 2) * 6.05; y = 1.7 + (index // 2) * 2.35
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.55), Inches(1.9))
            card.fill.solid(); card.fill.fore_color.rgb = _rgb("FFFFFF"); card.line.color.rgb = _rgb("DDE5EF")
            _textbox(slide, x + .3, y + .22, .55, .4, str(index + 1).zfill(2), 13, theme["accent"], True)
            _textbox(slide, x + .3, y + .68, 4.95, .92, value, 18 if len(value) > 45 else 21, theme["text"], True)

    def _image(self, slide, item, theme):
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.15), Inches(1.65), Inches(6.25), Inches(4.95))
        panel.fill.solid(); panel.fill.fore_color.rgb = _rgb("FFFFFF"); panel.line.color.rgb = _rgb("DDE5EF")
        slide.shapes.add_picture(item.image_path, Inches(6.28), Inches(1.78), width=Inches(5.99), height=Inches(4.69))
        for index, value in enumerate(item.bullets[:4]):
            _textbox(slide, .85, 1.75 + index * 1.05, 4.9, .82, "•  " + value, 17 if len(value) > 45 else 19, theme["text"])

    def _section(self, slide, item, theme):
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.8), Inches(1.75), Inches(11.75), Inches(4.65))
        panel.fill.solid(); panel.fill.fore_color.rgb = _rgb(theme["dark"]); panel.line.fill.background()
        _textbox(slide, 1.35, 2.35, 10.65, 1.25, item.subtitle or item.title, 34, "FFFFFF", True, PP_ALIGN.CENTER)
        if item.bullets:
            _textbox(slide, 2.0, 4.05, 9.3, 1.0, " · ".join(item.bullets[:3]), 19, theme["accent"], False, PP_ALIGN.CENTER)

    def _summary(self, slide, item, theme):
        for i, value in enumerate(item.bullets[:3]):
            x = .8 + i * 4.15
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.05), Inches(3.75), Inches(2.7))
            shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(theme["dark"]); shape.line.fill.background()
            _textbox(slide, x + .25, 2.35, 3.2, .45, str(i + 1).zfill(2), 14, theme["accent"], True)
            _textbox(slide, x + .28, 3.0, 3.15, 1.1, value, 22, "FFFFFF", True, PP_ALIGN.CENTER)

    def _chart(self, slide, item, theme):
        data = ChartData(); data.categories = item.chart.categories; data.add_series(item.chart.series_name, item.chart.values)
        kind = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}[item.chart.type]
        chart_x, text_x = ((5.05, .85) if item.layout == "chart-right-text-left" else (.8, 8.75))
        chart = slide.shapes.add_chart(kind, Inches(chart_x), Inches(1.65), Inches(7.5), Inches(4.9), data).chart
        chart.has_legend = True; chart.legend.position = XL_LEGEND_POSITION.BOTTOM; chart.has_title = False
        if item.chart.type == "pie":
            plot = chart.plots[0]; plot.has_data_labels = True
            plot.data_labels.position = XL_LABEL_POSITION.BEST_FIT
            plot.data_labels.show_category_name = True; plot.data_labels.show_percentage = True
        else:
            plot = chart.plots[0]; plot.has_data_labels = True
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            plot.data_labels.show_value = True
        _textbox(slide, text_x, 1.9, 3.7, .5, "关键结论", 18, theme["primary"], True)
        for i, value in enumerate(item.bullets[:4]): _textbox(slide, text_x, 2.6 + i * .85, 3.7, .66, "•  " + value, 16, theme["text"])
